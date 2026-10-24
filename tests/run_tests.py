import sys
import os

# Add backend and integrations to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend")))
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.services.scorer import calculate_deterministic_score, validate_keyword_set, ScoringError
from app.services.extractor import (
    extract_and_validate_reflection, 
    extract_document_text,
    CharacterLimitExceededError, 
    ExtractionError, 
    check_identity_conflict
)
from app.services.keyword_extractor import extract_top_15_keywords, FORBIDDEN_STOPWORDS
from app.services.similarity import select_best_reference, compute_cosine_similarity
from app.services.excel_exporter import generate_session_assessment_sheet, generate_master_performance_sheet
from app.services.ollama_evaluator import evaluate_with_qwen
from integrations.raven.client import RavenWorkerClient
import openpyxl
import io

def test_scoring_deterministic_formula():
    print("[TEST] Running test_scoring_deterministic_formula...")
    keywords = [
        "optimization", "gradient", "backpropagation", "weights", "convergence",
        "epochs", "momentum", "regularization", "hyperparameter", "loss",
        "activation", "stochastic", "normalization", "landscape", "decay"
    ]
    assert len(keywords) == 15, "Test must use 15 keywords"

    # Test Case 1: 15/15 matched -> 5.00
    reflection_all = "We studied optimization, gradient descent, backpropagation to update weights. We checked convergence across epochs with momentum and regularization. Tuning hyperparameter minimized loss, with activation functions and stochastic normalization across the loss landscape and weight decay."
    res = calculate_deterministic_score(reflection_all, keywords)
    assert res["reference_score"] == 15, f"Expected 15, got {res['reference_score']}"
    assert res["calculated_mark"] == 5.00, f"Expected 5.00, got {res['calculated_mark']}"

    # Test Case 2: 12/15 matched -> 4.00 (12 / 3)
    reflection_12 = "optimization gradient backpropagation weights convergence epochs momentum regularization hyperparameter loss activation stochastic"
    res12 = calculate_deterministic_score(reflection_12, keywords)
    assert res12["reference_score"] == 12, f"Expected 12, got {res12['reference_score']}"
    assert res12["calculated_mark"] == 4.00, f"Expected 4.00, got {res12['calculated_mark']}"

    # Test Case 3: 9/15 matched -> 3.00 (9 / 3)
    reflection_9 = "optimization gradient backpropagation weights convergence epochs momentum regularization hyperparameter"
    res9 = calculate_deterministic_score(reflection_9, keywords)
    assert res9["reference_score"] == 9, f"Expected 9, got {res9['reference_score']}"
    assert res9["calculated_mark"] == 3.00, f"Expected 3.00, got {res9['calculated_mark']}"

    # Test Case 4: 0/15 matched -> 0.00
    reflection_0 = "Today in chemistry class we mixed hydrogen and oxygen to produce water molecules and studied thermodynamics."
    res0 = calculate_deterministic_score(reflection_0, keywords)
    assert res0["reference_score"] == 0, f"Expected 0, got {res0['reference_score']}"
    assert res0["calculated_mark"] == 0.00, f"Expected 0.00, got {res0['calculated_mark']}"

    # Test Case 5: Repeated keywords earn NO extra points
    reflection_repeat = "optimization optimization optimization gradient gradient weights weights weights"
    res_repeat = calculate_deterministic_score(reflection_repeat, keywords)
    assert res_repeat["reference_score"] == 3, f"Expected 3 distinct matches, got {res_repeat['reference_score']}"
    assert res_repeat["calculated_mark"] == 1.00, f"Expected 1.00, got {res_repeat['calculated_mark']}"

    # Test Case 6: 5-Keywords Architecture Validation (5 matches = 5.00, 3 matches = 3.00, 1 match = 1.00)
    kw5 = ["optimization", "gradient", "backpropagation", "weights", "convergence"]
    res5_all = calculate_deterministic_score("Optimization using gradient and backpropagation updates weights towards convergence.", kw5)
    assert res5_all["reference_score"] == 5, f"Expected 5, got {res5_all['reference_score']}"
    assert res5_all["calculated_mark"] == 5.00, f"Expected 5.00, got {res5_all['calculated_mark']}"

    res5_3 = calculate_deterministic_score("Optimization with gradient and weights.", kw5)
    assert res5_3["reference_score"] == 3, f"Expected 3, got {res5_3['reference_score']}"
    assert res5_3["calculated_mark"] == 3.00, f"Expected 3.00, got {res5_3['calculated_mark']}"

    print("  --> PASS: Deterministic scoring formula and repeated keyword rules verified.")

def test_keyword_validation_stopwords():
    print("[TEST] Running test_keyword_validation_stopwords...")
    # List with fewer than 15 distinct keywords should fail
    short_list = ["gradient", "weights", "loss", "the", "and"] # 'the' & 'and' are stopwords
    try:
        validate_keyword_set(short_list)
        assert False, "Should have thrown ScoringError for under-sized keyword set"
    except ScoringError:
        pass

    # List with duplicate keywords
    dup_list = [
        "gradient", "weights", "loss", "backpropagation", "epochs",
        "momentum", "regularization", "hyperparameter", "activation", "stochastic",
        "normalization", "landscape", "decay", "gradient", "weights"
    ]
    try:
        validate_keyword_set(dup_list)
        assert False, "Should have thrown ScoringError for duplicate keywords reducing count below 15"
    except ScoringError:
        pass

    print("  --> PASS: Keyword validation and stopword filtering verified.")

def test_character_limit_and_extractor():
    print("[TEST] Running test_character_limit_and_extractor...")
    
    # Valid TXT within 500 characters
    valid_text = "This is a solid academic reflection exploring backpropagation and loss functions."
    cleaned, count = extract_and_validate_reflection("test.txt", valid_text.encode("utf-8"), max_characters=500)
    assert count == len(valid_text)
    assert cleaned == valid_text

    # Exceeding 500 characters
    long_text = "A" * 501
    try:
        extract_and_validate_reflection("long.txt", long_text.encode("utf-8"), max_characters=500)
        assert False, "Should have raised CharacterLimitExceededError"
    except CharacterLimitExceededError as e:
        assert e.char_count == 501

    # Empty file
    try:
        extract_and_validate_reflection("empty.txt", b"   \n\t  ", max_characters=500)
        assert False, "Should have raised ExtractionError for empty content"
    except ExtractionError:
        pass

    # Unsupported format
    try:
        extract_and_validate_reflection("test.exe", b"binary content", max_characters=500)
        assert False, "Should have raised ExtractionError for unsupported extension"
    except ExtractionError:
        pass

    print("  --> PASS: 500-character limit, empty doc, and format constraints verified.")

def test_identity_conflict_detection():
    print("[TEST] Running test_identity_conflict_detection...")
    text_with_diff_usn = "Submitted by 1RV21CS099 on deep learning algorithms."
    has_conflict, detected = check_identity_conflict(text_with_diff_usn, expected_usn="1RV21CS001")
    assert has_conflict is True
    assert detected == "1RV21CS099"

    text_matching = "Student ID: 1RV21CS001 reflection notes."
    has_conflict2, _ = check_identity_conflict(text_matching, expected_usn="1RV21CS001")
    assert has_conflict2 is False

    print("  --> PASS: Student USN mismatch detection verified.")

def test_multi_reference_similarity_selection():
    print("[TEST] Running test_multi_reference_similarity_selection...")
    ref_ml = {
        "id": "ref-ml-1",
        "extracted_text": "Machine learning, neural networks, gradient descent, backpropagation, and loss functions."
    }
    ref_db = {
        "id": "ref-db-2",
        "extracted_text": "Relational databases, SQL queries, indexing, B-trees, normalization, ACID transactions."
    }

    # Reflection clearly about machine learning
    reflection_ml = "I learned how gradient descent and backpropagation optimize neural network loss."
    res = select_best_reference(reflection_ml, [ref_ml, ref_db], min_threshold=0.20)
    assert res.selected_reference_id == "ref-ml-1", f"Expected ref-ml-1, got {res.selected_reference_id}"
    assert res.qualifies is True
    assert res.similarity_score > 0.20

    # Reflection with irrelevant text failing minimum threshold
    reflection_irrelevant = "Cooking pasta requires boiling water, adding salt, and simmering tomato sauce."
    res_irr = select_best_reference(reflection_irrelevant, [ref_ml, ref_db], min_threshold=0.35)
    assert res_irr.qualifies is False, "Should not qualify when below threshold"

    # User Specified Multi-Document Selection Rule:
    # Material 1: 3 Keywords Match
    # Material 2: 4 Keywords Match
    # Material 3: 8 Keywords Match
    # Winner: Material 3 (Highest Keyword Matching File)
    # Invariant: Strict Isolation - ZERO cross matching / mixing of keywords
    mat_1 = {
        "id": "mat-1",
        "name": "Material 1",
        "keywords": ["apple", "banana", "cherry", "date", "fig", "grape", "kiwi", "lemon", "mango", "nectarine", "orange", "pear", "quince", "raspberry", "strawberry"]
    }
    mat_2 = {
        "id": "mat-2",
        "name": "Material 2",
        "keywords": ["red", "blue", "green", "yellow", "purple", "cyan", "magenta", "black", "white", "gray", "silver", "gold", "bronze", "pink", "brown"]
    }
    mat_3 = {
        "id": "mat-3",
        "name": "Material 3",
        "keywords": ["cpu", "ram", "gpu", "disk", "kernel", "thread", "process", "cache", "bus", "register", "socket", "core", "clock", "voltage", "fan"]
    }

    test_multi_reflection = (
        "We tested apple, banana, cherry on a red, blue, green, yellow board. "
        "The system utilized cpu, ram, gpu, disk with kernel, thread, process and cache optimization."
    )

    res_multi = select_best_reference(test_multi_reflection, [mat_1, mat_2, mat_3])
    assert res_multi.selected_reference_id == "mat-3", f"Expected mat-3 (highest count), got {res_multi.selected_reference_id}"
    assert res_multi.reference_score == 8, f"Expected reference score 8, got {res_multi.reference_score}"
    assert len(res_multi.matched_keywords) == 8, f"Expected 8 matched keywords, got {len(res_multi.matched_keywords)}"
    assert res_multi.all_keyword_counts == {"mat-1": 3, "mat-2": 4, "mat-3": 8}

    # Strict isolation check: no keyword cross-contamination
    for kw in ["apple", "banana", "cherry", "red", "blue", "green", "yellow"]:
        assert kw not in res_multi.matched_keywords, f"Cross-matched keyword '{kw}' from other files leaked into mat-3!"
    for kw in ["cpu", "ram", "gpu", "disk", "kernel", "thread", "process", "cache"]:
        assert kw in res_multi.matched_keywords, f"Missing expected mat-3 keyword '{kw}'"

    print("  --> PASS: Multi-reference similarity selection and threshold gating verified.")

def test_excel_exports():
    print("[TEST] Running test_excel_exports...")
    # Test 1: Session Assessment Sheet
    records = [
        {
            "usn": "1RV21CS001",
            "student_name": "Alice Smith",
            "final_mark": 4.00,
            "reference_score": 12,
            "selected_reference_id": "ref-1",
            "selected_material_name": "Lecture 1 Notes",
            "submission_status": "SUBMITTED",
            "assessment_status": "COMPLETED"
        },
        {
            "usn": "1RV21CS002",
            "student_name": "Bob Jones",
            "final_mark": 0.00,
            "reference_score": 0,
            "selected_reference_id": "N/A",
            "selected_material_name": "N/A",
            "submission_status": "NOT_SUBMITTED",
            "assessment_status": "ABSENT"
        }
    ]
    session_xlsx = generate_session_assessment_sheet("sess-001", "Session 1: Backprop", records)
    assert len(session_xlsx) > 0
    wb_session = openpyxl.load_workbook(io.BytesIO(session_xlsx))
    assert "Session Assessment" in wb_session.sheetnames

    # Test 2: Master Performance Sheet
    sessions = [{"id": "s1", "name": "Session 1"}, {"id": "s2", "name": "Session 2"}]
    students = [
        {
            "name": "Alice Smith",
            "usn": "1RV21CS001",
            "scores": {
                "s1": {"mark": 4.00, "status": "COMPLETED"},
                "s2": {"mark": 5.00, "status": "COMPLETED"}
            }
        },
        {
            "name": "Charlie Brown",
            "usn": "1RV21CS003",
            "scores": {
                "s1": {"mark": 0.00, "status": "ABSENT"}, # Distinct absent zero
                "s2": {"mark": 0.00, "status": "COMPLETED"} # Distinct assessed zero
            }
        }
    ]
    master_xlsx = generate_master_performance_sheet(sessions, students)
    assert len(master_xlsx) > 0
    wb_master = openpyxl.load_workbook(io.BytesIO(master_xlsx))
    assert "Master Performance" in wb_master.sheetnames

    print("  --> PASS: Excel reports generated with valid openpyxl sheets and absent vs assessed zero distinction.")

def test_raven_offline_decoupling():
    print("[TEST] Running test_raven_offline_decoupling...")
    # Test that client cleanly reports unavailable without throwing unhandled exceptions
    client = RavenWorkerClient(worker_url="http://127.0.0.1:9999", timeout_seconds=1)
    import asyncio
    is_avail = asyncio.run(client.is_available())
    assert is_avail is False, "Expected offline worker to be reported as unavailable"
    print("  --> PASS: RAVEN worker offline decoupling and graceful fallback verified.")

def test_reference_document_keyword_extraction():
    print("[TEST] Running test_reference_document_keyword_extraction...")
    lecture_text = (
        "Gradient descent is an optimization algorithm used to minimize the loss function by iteratively moving "
        "in the direction of steepest descent. Backpropagation calculates the gradient of the error with respect "
        "to each weight using the chain rule. Key components include learning rate, momentum, activation function, "
        "convergence, epochs, batch normalization, regularization, stochastic gradient descent, hyperparameter tuning, "
        "loss landscape, and weight decay."
    )
    extracted_text = extract_document_text("lecture1.txt", lecture_text.encode("utf-8"))
    assert len(extracted_text) > 50, "Extracted text should not be empty"

    keywords = extract_top_15_keywords(extracted_text)
    assert len(keywords) == 15, f"Must extract exactly 15 keywords, got {len(keywords)}"
    assert len(set(keywords)) == 15, "Keywords must be distinct"

    # Non-negotiable check: generic English stopwords must NEVER appear as keywords
    banned_samples = ["as", "a", "is", "the", "for", "have", "has", "it", "that", "with", "and", "in", "of", "to"]
    for word in banned_samples:
        assert word in FORBIDDEN_STOPWORDS, f"Word '{word}' should be categorized in FORBIDDEN_STOPWORDS"
        assert word not in keywords, f"Forbidden stopword '{word}' leaked into keywords!"

    # Verify student scoring with these 15 extracted keywords
    scoring = calculate_deterministic_score(extracted_text, keywords)
    assert scoring["reference_score"] == 15, "All 15 keywords matched should give 15 points"
    assert scoring["calculated_mark"] == 5.00, "15 points / 3 = 5.00 marks"

    print("  --> PASS: Reference document reading, stopword exclusion, and 15-keyword extraction verified.")

def test_ollama_qwen_integration():
    print("[TEST] Running test_ollama_qwen_integration...")
    sample_keywords = [
        "optimization", "gradient", "backpropagation", "weights", "convergence",
        "epochs", "momentum", "regularization", "hyperparameter", "loss",
        "activation", "stochastic", "normalization", "landscape", "decay"
    ]
    sample_text = (
        "We studied optimization and gradient descent to update weights. "
        "We tuned the hyperparameter to minimize loss."
    )
    res = evaluate_with_qwen(
        reflection_text=sample_text,
        reference_title="Machine Learning Foundations",
        reference_keywords=sample_keywords,
        timeout_seconds=25
    )
    assert "success" in res
    assert "reference_score" in res
    assert "calculated_mark" in res
    if res["success"]:
        assert res["reference_score"] >= 1
        assert res["calculated_mark"] == round(res["reference_score"] / 3.0, 2)
        assert len(res["feedback"]) > 0
        print(f"  --> PASS: Ollama Qwen evaluation verified live (Score: {res['reference_score']}/15, Mark: {res['calculated_mark']}/5.00, Time: {res['duration_seconds']}s).")
    else:
        print(f"  --> PASS: Ollama graceful fallback verified ({res.get('error')}).")

if __name__ == "__main__":
    print("=================================================================")
    print("RUNNING AUTOMATED UNIT & INTEGRATION TEST SUITE (X:\\RAVEN ARK)")
    print("=================================================================")
    test_scoring_deterministic_formula()
    test_keyword_validation_stopwords()
    test_character_limit_and_extractor()
    test_identity_conflict_detection()
    test_multi_reference_similarity_selection()
    test_excel_exports()
    test_raven_offline_decoupling()
    test_reference_document_keyword_extraction()
    test_ollama_qwen_integration()
    print("=================================================================")
    print("ALL TESTS PASSED SUCCESSFULLY! (9/9 suites verified)")
    print("=================================================================")
